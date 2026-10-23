"""Word output: verified save, headings and Assumptions, labels, safe names (plan group 9)."""

import pytest
from docx import Document

from app.docx_builder import (
    ASSUMPTIONS_HEADING,
    DocxRenderer,
    RenderError,
    _save_verified,
    build_document,
    is_safe_document_name,
)
from app.schemas import Analysis, ColumnLabels, Item, Section, Table


@pytest.fixture
def plan_analysis() -> Analysis:
    return Analysis(
        is_document_request=True,
        title="CRM Migration Plan",
        document_type="project plan",
        category="planning",
        assumptions=[{"id": "A1", "statement": "Migration runs over 8 weeks",
                      "rationale": "No timeline was given."}],
    )


@pytest.fixture
def plan_sections() -> list[Section]:
    timeline = Table(
        columns=ColumnLabels(text="Milestone", owner="Owner", due="Target"),
        rows=[
            Item(type="task", kind="record", text="Data migration", owner="Priya", due="March 31"),
            Item(type="milestone", kind="proposal", text="Pilot with sales team", due="Week 5",
                 label="Proposed"),
        ],
    )
    return [
        Section(task_id="T1", heading="Overview", paragraphs=["This plan covers moving to the new CRM."]),
        Section(task_id="T2", heading="Timeline", table=timeline),
        Section(task_id="T3", heading="Risks",
                missing_source_note="No information for 'Risks' was provided in the request."),
    ]


def _headings(document) -> list[str]:
    return [p.text for p in document.paragraphs if p.style.name.startswith("Heading")]


def test_rendered_file_reopens_with_title_sections_and_assumptions(
    tmp_path, plan_analysis: Analysis, plan_sections: list[Section]
) -> None:
    name = DocxRenderer(tmp_path)(plan_analysis, plan_sections)

    document = Document(str(tmp_path / name))
    texts = [p.text for p in document.paragraphs]

    assert is_safe_document_name(name)
    assert name.endswith("-crm-migration-plan.docx")
    assert _headings(document) == ["Overview", "Timeline", "Risks", ASSUMPTIONS_HEADING]
    assert texts[0] == "CRM Migration Plan"
    assert any(text.startswith("Items marked Proposed") for text in texts)  # legend, as a proposal exists
    assert "Assumption A1: Migration runs over 8 weeks — No timeline was given." in texts
    rows = [[cell.text for cell in row.cells] for row in document.tables[0].rows]
    assert rows[0] == ["Milestone", "Owner", "Target", "Basis"]
    assert rows[1][-1] == "Recorded" and rows[2][-1] == "Proposed"


def test_failed_verification_leaves_no_file_behind(
    tmp_path, plan_analysis: Analysis, plan_sections: list[Section]
) -> None:
    document = build_document(plan_analysis, plan_sections)

    with pytest.raises(RenderError):
        _save_verified(document, tmp_path / "never.docx", plan_analysis.title, expected_headings=99)
    assert list(tmp_path.iterdir()) == []


@pytest.mark.parametrize("name", [
    "../secret.docx",
    "..\\x.docx",
    "a.docx",
    f"{'a' * 32}-ok.docx.exe",
    f"{'A' * 32}-x.docx",
])
def test_unsafe_document_names_are_rejected(name: str) -> None:
    assert not is_safe_document_name(name)


def test_value_that_repeats_the_owner_is_not_shown(plan_analysis: Analysis) -> None:
    # Seen live: a phase item rendered "Owner: Abhinav; ...; Value: Owner: Abhinav (Project Lead)".
    phase = Item(type="milestone", kind="proposal", label="Proposed", text="Phase 1",
                 owner="Abhinav", value="Owner: Abhinav (Project Lead)")
    table = Table(columns=ColumnLabels(text="Phase"), rows=[phase])
    sections = [Section(task_id="T1", heading="Phases", items=[phase]),
                Section(task_id="T2", heading="Phase table", table=table)]

    document = build_document(plan_analysis, sections)

    assert not any("Value" in paragraph.text for paragraph in document.paragraphs)
    assert [cell.text for cell in document.tables[0].rows[0].cells] == ["Phase", "Owner", "Basis"]


def test_table_hides_unused_columns_and_replaces_field_name_headers(plan_analysis: Analysis) -> None:
    # Seen live: the model labelled columns "text/owner/due/value" and the value column stayed empty.
    table = Table(
        columns=ColumnLabels(text="text", owner="owner", due="Deadline", value="value"),
        rows=[Item(type="action", kind="record", text="Send the client list", owner="Arjun", due="24 Oct")],
    )
    sections = [Section(task_id="T1", heading="Action Items", table=table)]

    document = build_document(plan_analysis, sections)

    header = [cell.text for cell in document.tables[0].rows[0].cells]
    assert header == ["Item", "Owner", "Deadline"]
